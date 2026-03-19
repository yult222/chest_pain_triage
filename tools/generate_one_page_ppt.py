from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(245, 247, 250)
NAVY = RGBColor(22, 41, 64)
TEXT = RGBColor(38, 50, 56)
MUTED = RGBColor(94, 109, 125)
RED = RGBColor(196, 48, 43)
RED_SOFT = RGBColor(255, 235, 233)
TEAL = RGBColor(17, 122, 101)
TEAL_SOFT = RGBColor(229, 247, 243)
BLUE = RGBColor(32, 84, 147)
BLUE_SOFT = RGBColor(233, 241, 252)
LINE = RGBColor(219, 226, 234)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(179, 118, 0)
GOLD_SOFT = RGBColor(255, 244, 219)

FONT = "Microsoft YaHei"
FONT_BOLD = "Microsoft YaHei"


def set_run_style(run, size: int, color: RGBColor, bold: bool = False) -> None:
    run.font.name = FONT_BOLD if bold else FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 18,
    color: RGBColor = TEXT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    margin_left=0.1,
    margin_right=0.1,
    margin_top=0.06,
    margin_bottom=0.04,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin_left)
    tf.margin_right = Inches(margin_right)
    tf.margin_top = Inches(margin_top)
    tf.margin_bottom = Inches(margin_bottom)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_style(run, size=size, color=color, bold=bold)
    return box


def add_round_rect(slide, left, top, width, height, fill: RGBColor, line_color: RGBColor = LINE, radius_type=None):
    shape_type = radius_type or MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line_color
    shape.line.width = Pt(1)
    return shape


def add_chip(slide, left, top, width, height, text: str, fill: RGBColor, color: RGBColor) -> None:
    shape = add_round_rect(slide, left, top, width, height, fill, fill)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    set_run_style(run, size=13, color=color, bold=True)


def add_card_title(slide, left, top, width, text: str) -> None:
    add_textbox(slide, left, top, width, Inches(0.28), text, size=13, color=MUTED, bold=True)


def add_flow_card(slide, left, top, width, height, *, title: str, body_lines: list[str], accent: RGBColor, soft: RGBColor):
    card = add_round_rect(slide, left, top, width, height, WHITE)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Inches(0.08), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.color.rgb = accent

    add_textbox(slide, left + Inches(0.18), top + Inches(0.16), width - Inches(0.28), Inches(0.35), title, size=19, color=NAVY, bold=True)
    badge = add_round_rect(slide, left + Inches(0.18), top + Inches(0.58), Inches(1.02), Inches(0.34), soft, soft)
    tf = badge.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "关键要点"
    set_run_style(run, size=11, color=accent, bold=True)

    current_top = top + Inches(1.03)
    for line in body_lines:
        add_textbox(
            slide,
            left + Inches(0.18),
            current_top,
            width - Inches(0.28),
            Inches(0.32),
            f"• {line}",
            size=13,
            color=TEXT,
            margin_left=0.0,
        )
        current_top += Inches(0.34)
    return card


def add_arrow(slide, left, top, width, height) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE_SOFT
    shape.line.color.rgb = BLUE_SOFT
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = ""


def add_bullet_panel(slide, left, top, width, height) -> None:
    panel = add_round_rect(slide, left, top, width, height, WHITE)
    add_card_title(slide, left + Inches(0.12), top + Inches(0.08), width - Inches(0.24), "三行看懂")

    rows = [
        ("输入", "胸痛部位、症状问卷、自由文本描述", BLUE, BLUE_SOFT),
        ("判断", "红旗征识别 + 医生可写规则引擎", RED, RED_SOFT),
        ("输出", "紧急度分级、推荐科室、病例摘要导出", TEAL, TEAL_SOFT),
    ]
    row_top = top + Inches(0.42)
    for label, text, color, fill in rows:
        chip = add_round_rect(slide, left + Inches(0.14), row_top + Inches(0.02), Inches(0.72), Inches(0.34), fill, fill)
        tf = chip.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        set_run_style(run, size=12, color=color, bold=True)

        add_textbox(
            slide,
            left + Inches(0.95),
            row_top,
            width - Inches(1.08),
            Inches(0.38),
            text,
            size=14,
            color=TEXT,
            margin_left=0.0,
        )
        row_top += Inches(0.49)
    return panel


def add_summary_panel(slide, left, top, width, height) -> None:
    panel = add_round_rect(slide, left, top, width, height, WHITE)
    add_card_title(slide, left + Inches(0.12), top + Inches(0.08), width - Inches(0.24), "一句话")
    summary = (
        "面向胸痛场景的分诊辅助工具，在不替代医生诊断的前提下，"
        "提高首轮分流效率与安全性。"
    )
    add_textbox(
        slide,
        left + Inches(0.12),
        top + Inches(0.38),
        width - Inches(0.24),
        height - Inches(0.48),
        summary,
        size=18,
        color=NAVY,
        bold=False,
        margin_left=0.0,
        margin_top=0.02,
    )
    return panel


def add_value_bar(slide, left, top, width, height) -> None:
    bar = add_round_rect(slide, left, top, width, height, WHITE)
    add_card_title(slide, left + Inches(0.12), top + Inches(0.06), width - Inches(0.24), "价值点")
    chips = [
        ("更快识别高危患者", RED_SOFT, RED),
        ("更清晰给出分诊依据", BLUE_SOFT, BLUE),
        ("更方便衔接后续接诊", TEAL_SOFT, TEAL),
    ]
    chip_left = left + Inches(0.14)
    chip_width = Inches(1.68)
    gap = Inches(0.14)
    for text, fill, color in chips:
        add_chip(slide, chip_left, top + Inches(0.36), chip_width, Inches(0.48), text, fill, color)
        chip_left += chip_width + gap
    return bar


def add_footer(slide) -> None:
    add_textbox(
        slide,
        Inches(0.7),
        Inches(6.98),
        Inches(12.0),
        Inches(0.22),
        "用于分诊辅助，不替代医生诊断；如出现严重胸痛、气促、晕厥等情况，应优先急诊评估。",
        size=10,
        color=MUTED,
        align=PP_ALIGN.CENTER,
        margin_left=0.0,
        margin_right=0.0,
        margin_top=0.0,
        margin_bottom=0.0,
    )


def add_background(slide) -> None:
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    blob1 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(-0.2), Inches(-0.5), Inches(2.6), Inches(2.6))
    blob1.fill.solid()
    blob1.fill.fore_color.rgb = RED_SOFT
    blob1.line.color.rgb = RED_SOFT

    blob2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(11.4), Inches(5.4), Inches(2.2), Inches(2.2))
    blob2.fill.solid()
    blob2.fill.fore_color.rgb = BLUE_SOFT
    blob2.line.color.rgb = BLUE_SOFT


def add_title_block(slide) -> None:
    add_chip(slide, Inches(0.7), Inches(0.48), Inches(1.52), Inches(0.34), "胸痛分诊辅助", RED_SOFT, RED)
    add_textbox(
        slide,
        Inches(0.7),
        Inches(0.88),
        Inches(6.9),
        Inches(0.56),
        "胸痛智能分诊辅助系统",
        size=26,
        color=NAVY,
        bold=True,
        margin_left=0.0,
        margin_top=0.0,
    )
    subtitle = "用结构化问诊与可解释规则，帮助更快识别高风险胸痛并推荐就诊路径"
    add_textbox(
        slide,
        Inches(0.72),
        Inches(1.42),
        Inches(7.2),
        Inches(0.34),
        subtitle,
        size=14,
        color=MUTED,
        margin_left=0.0,
        margin_top=0.0,
    )


def add_visual_panel(slide, screenshot: Path | None) -> None:
    panel_left = Inches(0.7)
    panel_top = Inches(2.0)
    panel_w = Inches(7.2)
    panel_h = Inches(4.45)
    panel = add_round_rect(slide, panel_left, panel_top, panel_w, panel_h, WHITE)
    add_card_title(slide, panel_left + Inches(0.12), panel_top + Inches(0.08), panel_w - Inches(0.24), "核心流程")

    if screenshot and screenshot.exists():
        pic_area_top = panel_top + Inches(0.4)
        slide.shapes.add_picture(str(screenshot), panel_left + Inches(0.18), pic_area_top, width=panel_w - Inches(0.36))
        overlay = add_round_rect(
            slide,
            panel_left + Inches(4.84),
            panel_top + Inches(0.22),
            Inches(2.0),
            Inches(0.38),
            RED_SOFT,
            RED_SOFT,
        )
        tf = overlay.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = "建议保留红旗提示区域"
        set_run_style(run, size=11, color=RED, bold=True)
        return panel

    card_top = panel_top + Inches(0.72)
    card_w = Inches(2.06)
    card_h = Inches(2.42)
    add_flow_card(
        slide,
        panel_left + Inches(0.18),
        card_top,
        card_w,
        card_h,
        title="采集信息",
        body_lines=["疼痛部位定位", "症状问卷补充", "自由文本描述"],
        accent=BLUE,
        soft=BLUE_SOFT,
    )
    add_arrow(slide, panel_left + Inches(2.42), card_top + Inches(0.82), Inches(0.42), Inches(0.48))
    add_flow_card(
        slide,
        panel_left + Inches(2.9),
        card_top,
        card_w,
        card_h,
        title="规则判断",
        body_lines=["红旗征预警", "先验 + 证据推断", "关键缺失项提醒"],
        accent=RED,
        soft=RED_SOFT,
    )
    add_arrow(slide, panel_left + Inches(5.14), card_top + Inches(0.82), Inches(0.42), Inches(0.48))
    add_flow_card(
        slide,
        panel_left + Inches(5.62),
        card_top,
        card_w,
        card_h,
        title="分诊输出",
        body_lines=["紧急度分级", "推荐就诊科室", "病例摘要导出"],
        accent=TEAL,
        soft=TEAL_SOFT,
    )

    alert = add_round_rect(slide, panel_left + Inches(0.24), panel_top + Inches(3.48), panel_w - Inches(0.48), Inches(0.7), RED_SOFT, RED_SOFT)
    tf = alert.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "安全边界：系统用于辅助分诊，不替代医生诊断；发现高危信号时优先触发急诊评估。"
    set_run_style(run, size=13, color=RED, bold=True)
    return panel


def build_ppt(output_path: Path, screenshot: Path | None = None) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_background(slide)
    add_title_block(slide)
    add_visual_panel(slide, screenshot)

    right_left = Inches(8.2)
    add_summary_panel(slide, right_left, Inches(2.0), Inches(4.45), Inches(1.25))
    add_bullet_panel(slide, right_left, Inches(3.42), Inches(4.45), Inches(1.95))
    add_value_bar(slide, right_left, Inches(5.56), Inches(4.45), Inches(1.14))

    add_textbox(
        slide,
        Inches(10.9),
        Inches(0.56),
        Inches(1.7),
        Inches(0.5),
        "评审版一页概览",
        size=12,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        margin_left=0.0,
        margin_right=0.0,
        margin_top=0.0,
    )
    tag = add_round_rect(slide, Inches(10.7), Inches(0.92), Inches(1.96), Inches(0.36), GOLD_SOFT, GOLD_SOFT)
    tf = tag.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "适合综合评审场景"
    set_run_style(run, size=11, color=GOLD, bold=True)

    add_footer(slide)
    prs.save(str(output_path))


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate a one-page PPT for the chest pain triage product.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("deliverables/chest_pain_one_page_overview.pptx"),
        help="Where to write the PPTX file.",
    )
    parser.add_argument(
        "--screenshot",
        type=Path,
        default=None,
        help="Optional app screenshot to embed in the visual panel.",
    )
    args = parser.parse_args()

    build_ppt(args.output, args.screenshot)
    print(f"Generated: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
